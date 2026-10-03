from datetime import datetime
from io import BytesIO

from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def clean_xml_text(value: str) -> str:
    if not value:
        return ""

    cleaned_characters = []

    for character in str(value):
        codepoint = ord(character)

        is_allowed_control = character in ("\n", "\t", "\r")
        is_valid_xml_character = (
            codepoint == 0x9
            or codepoint == 0xA
            or codepoint == 0xD
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        )

        if is_allowed_control or is_valid_xml_character:
            cleaned_characters.append(character)

    return "".join(cleaned_characters)

def safe_filename(value: str) -> str:
    if not value:
        return "translated-document"

    safe_value = ""

    for character in value.lower():
        if character.isalnum():
            safe_value += character
        elif character in (" ", "-", "_"):
            safe_value += "-"

    while "--" in safe_value:
        safe_value = safe_value.replace("--", "-")

    safe_value = safe_value.strip("-")

    return safe_value[:80] or "translated-document"

def set_paragraph_rtl(paragraph) -> None:
    paragraph_format = paragraph.paragraph_format
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    paragraph_properties = paragraph._p.get_or_add_pPr()

    bidi = paragraph_properties.find(qn("w:bidi"))

    if bidi is None:
        bidi = OxmlElement("w:bidi")
        paragraph_properties.append(bidi)

    bidi.set(qn("w:val"), "1")


def set_run_rtl(run) -> None:
    run_properties = run._r.get_or_add_rPr()

    rtl = run_properties.find(qn("w:rtl"))

    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run_properties.append(rtl)

    rtl.set(qn("w:val"), "1")


def set_document_styles(document: Document) -> None:
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(13)

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(22)
    title_style.font.bold = True

    heading_1_style = styles["Heading 1"]
    heading_1_style.font.name = "Arial"
    heading_1_style.font.size = Pt(18)
    heading_1_style.font.bold = True

    heading_2_style = styles["Heading 2"]
    heading_2_style.font.name = "Arial"
    heading_2_style.font.size = Pt(15)
    heading_2_style.font.bold = True


def set_document_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def add_footer(section, source_type: str, source_number: int, review_status: str) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]

    exported_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    paragraph.text = (
        f"Smart Study Assistant & Uyghur Translator | "
        f"{source_type.title()} {source_number} | "
        f"Status: {review_status} | "
        f"Exported: {exported_at}"
    )

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.name = "Arial"


def add_ltr_metadata_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.name = "Arial"
    label_run.font.size = Pt(10)

    value_run = paragraph.add_run(clean_xml_text(value))
    value_run.font.name = "Arial"
    value_run.font.size = Pt(10)


def add_rtl_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(level=level)
    set_paragraph_rtl(paragraph)

    run = paragraph.add_run(clean_xml_text(text))
    run.font.name = "Arial"
    run.font.size = Pt(18 if level == 1 else 15)
    run.bold = True
    set_run_rtl(run)


def add_rtl_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_rtl(paragraph)

    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(6)
    paragraph_format.line_spacing = 1.25
    paragraph_format.left_indent = Inches(0)
    paragraph_format.right_indent = Inches(0)
    paragraph_format.first_line_indent = Inches(0)

    run = paragraph.add_run(clean_xml_text(text))
    run.font.name = "Arial"
    run.font.size = Pt(13)
    set_run_rtl(run)


def add_translation_body(document: Document, translated_text: str) -> None:
    cleaned_text = clean_xml_text(translated_text)
    paragraphs = cleaned_text.split("\n")

    for paragraph_text in paragraphs:
        stripped_text = paragraph_text.strip()

        if not stripped_text:
            document.add_paragraph("")
            continue

        add_rtl_paragraph(document, stripped_text)


def build_docx(
    *,
    title: str,
    original_text: str,
    translated_text: str,
    source_type: str,
    source_number: int,
    review_status: str,
    glossary_terms: str = "",
) -> BytesIO:
    document = Document()

    set_document_styles(document)
    set_document_margins(document)
    add_footer(document.sections[0], source_type, source_number, review_status)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(clean_xml_text(title))
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)

    title_paragraph.paragraph_format.space_after = Pt(18)
    add_translation_body(document, translated_text)


    if glossary_terms.strip():
        document.add_page_break()
        add_rtl_heading(document, "ئاتالغۇلار", level=1)

        for line in glossary_terms.splitlines():
            stripped_line = line.strip()

            if stripped_line:
                add_rtl_paragraph(document, stripped_line)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output

def build_combined_pages_docx(
    *,
    title: str,
    pages: list[dict],
    review_status: str = "translated",
    glossary_terms: str = "",
) -> BytesIO:
    document = Document()

    set_document_styles(document)
    set_document_margins(document)
    add_footer(document.sections[0], "combined pages", 0, review_status)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(18)

    title_run = title_paragraph.add_run(clean_xml_text(title))
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)

    sorted_pages = sorted(
        pages,
        key=lambda page: int(page.get("page_number", 0)),
    )

    for index, page in enumerate(sorted_pages):
        if index > 0:
            document.add_page_break()

        page_number = page.get("page_number", index + 1)
        page_title = page.get("title") or f"Page {page_number}"
        translated_text = page.get("translated_text", "")

        heading_paragraph = document.add_paragraph()
        heading_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_paragraph.paragraph_format.space_after = Pt(12)

        heading_run = heading_paragraph.add_run(clean_xml_text(page_title))
        heading_run.bold = True
        heading_run.font.name = "Arial"
        heading_run.font.size = Pt(15)

        add_translation_body(document, translated_text)

    if glossary_terms.strip():
        document.add_page_break()
        add_rtl_heading(document, "ئاتالغۇلار", level=1)

        for line in glossary_terms.splitlines():
            stripped_line = line.strip()

            if stripped_line:
                add_rtl_paragraph(document, stripped_line)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output